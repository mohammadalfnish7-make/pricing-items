"""Parse order item files and compute per-line costs."""

from __future__ import annotations

import csv
import io
import re
from typing import Any

import xlrd
from openpyxl import load_workbook

from app.catalog import enrich_item_from_catalog, load_catalog

# Accept both English snake_case and original bilingual headers
HEADER_ALIASES: dict[str, str] = {
    "order_id": "order_id",
    "line_no": "line_no",
    "no. / №": "line_no",
    "product_name": "product_name",
    "name of product / наименование товара": "product_name",
    "product_description": "product_description",
    "спецописание товара": "product_description",
    "its_rate": "its_rate",
    "итс": "its_rate",
    "total_amount": "total_amount",
    "итого / ставка": "total_amount",
    "итого": "total_amount",
    "ставка": "total_amount",
    "manufacturer": "manufacturer",
    "manufacturer / производитель": "manufacturer",
    "brand": "brand",
    "brand (trade mark) / бренд (товарный знак)": "brand",
    "tariff_code": "tariff_code",
    "country_of_origin": "country_of_origin",
    "country of origin / страна происхождения": "country_of_origin",
    "unit": "unit",
    "unit / ед. изм.": "unit",
    "quantity_units": "quantity_units",
    "q-ty, ps / кол-во, ед": "quantity_units",
    "quantity_packages": "quantity_packages",
    "q-ty, pks / кол-во, мест": "quantity_packages",
    "gross_weight_kg": "gross_weight_kg",
    "gross weight, kg / вес брутто, кг": "gross_weight_kg",
    "net_weight_kg": "net_weight_kg",
    "net weight, kg / вес нетто, кг": "net_weight_kg",
    "vat_rate": "vat_rate",
    "ставка ндс": "vat_rate",
    "price": "price",
    "цена": "price",
}


def _norm_header(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).replace("\n", " ")
    text = re.sub(r"\s+", " ", text).strip().lower()
    return text


def _to_float(value: Any) -> float | None:
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace(" ", "").replace(",", ".")
    if not text:
        return None
    try:
        return float(text)
    except ValueError:
        return None


def parse_its_rate(value: Any) -> tuple[float | None, str | None]:
    """Parse strings like '1,85 USD/кг' -> (1.85, 'USD/kg')."""
    if value is None or str(value).strip() == "":
        return None, None
    text = str(value).strip()
    match = re.match(
        r"^([\d]+(?:[.,]\d+)?)\s*([A-Za-zА-Яа-я$€£]+)?\s*/?\s*(.*)?$",
        text,
    )
    if not match:
        return None, text
    amount = _to_float(match.group(1))
    currency = (match.group(2) or "USD").strip()
    unit = (match.group(3) or "kg").strip() or "kg"
    unit = unit.replace("кг", "kg")
    return amount, f"{currency}/{unit}"


def map_row(raw: dict[str, Any]) -> dict[str, Any]:
    mapped: dict[str, Any] = {}
    for key, value in raw.items():
        canon = HEADER_ALIASES.get(_norm_header(key))
        if canon and canon not in mapped:
            mapped[canon] = value
    return mapped


def compute_item_cost(mapped: dict[str, Any]) -> dict[str, Any]:
    """Actual $/kg = Итого ÷ gross weight.

    Order Excel: E (Итого) ÷ N (gross kg)
    Merged CSV:  F (total_amount) ÷ Q (gross_weight_kg) — not E (ITS)

    line_cost = Итого
    price_per_kg = Итого / gross_weight_kg
    kg_per_unit (gross) = gross_weight / quantity
    cost_per_unit = Итого / quantity
    """
    itogo = _to_float(mapped.get("total_amount"))  # Итого
    quantity = _to_float(mapped.get("quantity_units"))
    net_weight = _to_float(mapped.get("net_weight_kg"))
    gross_weight = _to_float(mapped.get("gross_weight_kg"))
    vat_rate = _to_float(mapped.get("vat_rate"))
    price = _to_float(mapped.get("price"))
    its_amount, its_unit = parse_its_rate(mapped.get("its_rate"))

    # Actual commercial $/kg from the order file
    price_per_kg = None
    if itogo is not None and gross_weight not in (None, 0):
        price_per_kg = round(itogo / gross_weight, 6)

    kg_per_unit = None
    if gross_weight is not None and quantity not in (None, 0):
        kg_per_unit = round(gross_weight / quantity, 6)

    line_cost = None
    cost_source = None
    if itogo is not None:
        line_cost = round(itogo, 2)
        cost_source = "itogo"
    elif price_per_kg is not None and gross_weight is not None:
        line_cost = round(price_per_kg * gross_weight, 2)
        cost_source = "usd_per_kg_x_gross"

    cost_per_unit = None
    if line_cost is not None and quantity not in (None, 0):
        cost_per_unit = round(line_cost / quantity, 4)
    elif price_per_kg is not None and kg_per_unit is not None:
        cost_per_unit = round(price_per_kg * kg_per_unit, 4)

    # Keep ITS as reference only (customs rate), not used for line pricing
    its_cost = None
    if its_amount is not None and net_weight is not None:
        its_cost = round(its_amount * net_weight, 2)

    vat_amount = None
    if line_cost is not None and vat_rate is not None:
        vat_amount = round(line_cost * vat_rate / 100.0, 2)

    return {
        "order_id": (str(mapped["order_id"]).strip() if mapped.get("order_id") else None),
        "line_no": (str(mapped["line_no"]).strip() if mapped.get("line_no") else None),
        "product_name": (
            str(mapped["product_name"]).strip() if mapped.get("product_name") else None
        ),
        "product_description": (
            str(mapped["product_description"]).strip()
            if mapped.get("product_description")
            else None
        ),
        "manufacturer": (
            str(mapped["manufacturer"]).strip() if mapped.get("manufacturer") else None
        ),
        "brand": str(mapped["brand"]).strip() if mapped.get("brand") else None,
        "tariff_code": (
            str(mapped["tariff_code"]).strip() if mapped.get("tariff_code") else None
        ),
        "country_of_origin": (
            str(mapped["country_of_origin"]).strip()
            if mapped.get("country_of_origin")
            else None
        ),
        "unit": str(mapped["unit"]).strip() if mapped.get("unit") else None,
        "quantity_units": quantity,
        "quantity_packages": _to_float(mapped.get("quantity_packages")),
        "gross_weight_kg": gross_weight,
        "net_weight_kg": net_weight,
        "kg_per_unit": kg_per_unit,
        "its_rate_raw": (
            str(mapped["its_rate"]).strip() if mapped.get("its_rate") else None
        ),
        "its_rate_amount": its_amount,
        "its_rate_unit": its_unit or ("USD/kg" if its_amount is not None else None),
        "price_per_kg": price_per_kg,
        "its_cost": its_cost,
        "declared_amount": itogo,
        "itogo": itogo,
        "vat_rate": vat_rate,
        "vat_amount": vat_amount,
        "unit_cost": cost_per_unit,
        "line_cost": line_cost,
        "cost_source": cost_source,
        "currency": "USD",
        "price": price,
        "pricing_note": "usd_per_kg = itogo / gross_weight_kg",
    }


def _rows_from_csv(content: bytes) -> list[dict[str, Any]]:
    text = content.decode("utf-8-sig")
    reader = csv.DictReader(io.StringIO(text))
    return [dict(row) for row in reader]


def _rows_from_matrix(matrix: list[tuple[Any, ...]]) -> list[dict[str, Any]]:
    """Build dict rows from a 2D matrix; supports headered sheets or name-only lists."""
    if not matrix:
        return []

    first = matrix[0]
    nonempty_first = [c for c in first if c is not None and str(c).strip() != ""]
    # Single-column name list (e.g. N 145784.xls) — no header row
    if len(nonempty_first) == 1 and len(first) == 1:
        out: list[dict[str, Any]] = []
        for i, row in enumerate(matrix, 1):
            name = row[0] if row else None
            if name is None or str(name).strip() == "":
                continue
            out.append({"line_no": i, "product_name": str(name).strip()})
        return out

    keys = [
        str(h).strip() if h is not None and str(h).strip() != "" else f"col_{i}"
        for i, h in enumerate(first)
    ]
    # If first row does not look like known headers and is a single product-like cell, treat as data
    mapped_keys = [HEADER_ALIASES.get(_norm_header(k)) for k in keys]
    if not any(mapped_keys) and len(nonempty_first) == 1:
        out = []
        for i, row in enumerate(matrix, 1):
            vals = [c for c in row if c is not None and str(c).strip() != ""]
            if not vals:
                continue
            out.append({"line_no": i, "product_name": str(vals[0]).strip()})
        return out

    out = []
    for row in matrix[1:]:
        if all(c is None or str(c).strip() == "" for c in row):
            continue
        out.append({keys[i]: (row[i] if i < len(row) else None) for i in range(len(keys))})
    return out


def _rows_from_xlsx(content: bytes) -> list[dict[str, Any]]:
    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    ws = wb.active
    matrix = [tuple(row) for row in ws.iter_rows(values_only=True)]
    wb.close()
    return _rows_from_matrix(matrix)


def _rows_from_xls(content: bytes) -> list[dict[str, Any]]:
    wb = xlrd.open_workbook(file_contents=content)
    sh = wb.sheet_by_index(0)
    matrix = [
        tuple(sh.cell_value(r, c) for c in range(sh.ncols)) for r in range(sh.nrows)
    ]
    return _rows_from_matrix(matrix)


def _summarize_items(
    filename: str,
    items: list[dict[str, Any]],
    skipped: int,
    match_mode: str,
) -> dict[str, Any]:
    priced = [i for i in items if i["line_cost"] is not None]
    matched = sum(1 for i in items if i.get("match_status") == "matched")
    unmatched = sum(1 for i in items if i.get("match_status") == "unmatched")
    country_diffs = sum(1 for i in items if i.get("country_diff"))
    total_cost = round(sum(i["line_cost"] for i in priced), 2)
    total_its = round(sum(i["its_cost"] for i in items if i["its_cost"] is not None), 2)
    total_vat = round(sum(i["vat_amount"] for i in items if i["vat_amount"] is not None), 2)
    return {
        "filename": filename,
        "match_mode": match_mode,
        "item_count": len(items),
        "priced_count": len(priced),
        "catalog_matched": matched,
        "catalog_unmatched": unmatched,
        "country_diff_count": country_diffs,
        "catalog_size": len(load_catalog()),
        "skipped_rows": skipped,
        "currency": "USD",
        "pricing_mode": "itogo_div_gross_kg",
        "totals": {
            "line_cost": total_cost,
            "its_cost": total_its,
            "vat_amount": total_vat,
        },
        "items": items,
    }


def process_upload(filename: str, content: bytes) -> dict[str, Any]:
    import copy

    lower = filename.lower()
    if lower.endswith(".csv"):
        raw_rows = _rows_from_csv(content)
    elif lower.endswith(".xlsx") or lower.endswith(".xlsm"):
        raw_rows = _rows_from_xlsx(content)
    elif lower.endswith(".xls"):
        raw_rows = _rows_from_xls(content)
    else:
        raise ValueError("Unsupported file type. Upload .csv, .xlsx, or .xls")

    base_items: list[dict[str, Any]] = []
    skipped = 0
    for raw in raw_rows:
        mapped = map_row(raw)
        if not mapped.get("product_name") and raw.get("product_name"):
            mapped["product_name"] = raw["product_name"]
        if not mapped.get("line_no") and raw.get("line_no") is not None:
            mapped["line_no"] = raw["line_no"]
        if not mapped.get("product_name") and not mapped.get("line_no"):
            skipped += 1
            continue
        base_items.append(compute_item_cost(mapped))

    def with_mode(mode: str) -> list[dict[str, Any]]:
        out: list[dict[str, Any]] = []
        for item in base_items:
            cloned = copy.deepcopy(item)
            if cloned.get("line_cost") is None:
                cloned = enrich_item_from_catalog(cloned, mode=mode)
            else:
                cloned["match_mode"] = mode
            out.append(cloned)
        return out

    strict_items = with_mode("strict")
    fuzzy_items = with_mode("fuzzy")

    strict_page = _summarize_items(filename, strict_items, skipped, "strict")
    fuzzy_page = _summarize_items(filename, fuzzy_items, skipped, "fuzzy")

    # Top-level = page 1 (strict) for compatibility; pages for UI tabs
    return {
        **strict_page,
        "pages": {
            "strict": strict_page,
            "fuzzy": fuzzy_page,
        },
    }
